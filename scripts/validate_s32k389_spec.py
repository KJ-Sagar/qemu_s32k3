#!/usr/bin/env python3
"""Validate S32K389 board constants against the official NXP spreadsheets.

This is a focused regression check for the values that are directly exposed to
firmware and QEMU: the interrupt vectors and the primary MMIO base addresses.
It intentionally checks only the subset that is explicitly documented in the
S32K3xx memory/interrupt map workbooks so we do not silently invent values.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
HEADER = ROOT / "hw/arm/s32k389.h"
MEMORY_WORKBOOK = ROOT / "NXP Manuals" / "S32K3xx_memory_map.xlsx"
INTERRUPT_WORKBOOK = ROOT / "NXP Manuals" / "S32K3xx_interrupt_map.xlsx"


def read_header_macros() -> dict[str, int]:
    text = HEADER.read_text(encoding="utf-8")
    macros: dict[str, int] = {}
    for name, value in re.findall(
        r"^#define\s+((?:S32K3|S32K389)[A-Z0-9_]+)\s+(0x[0-9A-Fa-f]+|\d+)\b",
        text,
        re.MULTILINE,
    ):
        macros[name] = int(value, 0)
    return macros


def normalize_name(name: str) -> str:
    value = str(name).strip().lower()
    value = value.replace("_", " ")
    value = re.sub(r"[^a-z0-9 ]+", " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def build_irq_lookup() -> dict[tuple[str, str], int]:
    ws = load_workbook(INTERRUPT_WORKBOOK, read_only=True)["Interrupts"]
    lookup: dict[tuple[str, str], int] = {}
    for row in ws.iter_rows(values_only=True):
        if len(row) < 23 or row[22] != "YES":
            continue
        inst = str(row[6] or "").strip()
        req = str(row[7] or "").strip()
        irq = row[2]
        if not inst or not req or irq is None:
            continue
        try:
            irq_num = int(irq)
        except (TypeError, ValueError):
            continue
        lookup[(normalize_name(inst), normalize_name(req))] = irq_num
    return lookup


def build_base_lookup() -> dict[str, int]:
    ws = load_workbook(MEMORY_WORKBOOK, read_only=True)["Peripherals"]
    lookup: dict[str, int] = {}
    for row in ws.iter_rows(values_only=True):
        if len(row) < 23 or row[22] != "YES":
            continue
        inst = str(row[0] or "").strip()
        start = row[3]
        if not inst or start is None:
            continue
        try:
            lookup[normalize_name(inst)] = int(str(start), 16)
        except (TypeError, ValueError):
            continue
    return lookup


def validate_irq_macros(macros: dict[str, int], irq_lookup: dict[tuple[str, str], int]) -> list[str]:
    expected = {
        "S32K3_LPI2C0_IRQ": ("lpi2c 0", "lpi2c master interrupt"),
        "S32K3_LPI2C1_IRQ": ("lpi2c 1", "lpi2c master interrupt"),
        "S32K3_LPSPI0_IRQ": ("lpspi 0", "lpspi interrupt"),
        "S32K3_LPSPI1_IRQ": ("lpspi 1", "lpspi interrupt"),
        "S32K3_LPSPI2_IRQ": ("lpspi 2", "lpspi interrupt"),
        "S32K3_LPSPI3_IRQ": ("lpspi 3", "lpspi interrupt"),
        "S32K3_LPSPI4_IRQ": ("lpspi 4", "lpspi interrupt"),
        "S32K3_LPSPI5_IRQ": ("lpspi 5", "lpspi interrupt"),
        "S32K3_ADC0_IRQ": ("adc 0", "end of conversion"),
        "S32K3_ADC1_IRQ": ("adc 1", "end of conversion"),
        "S32K3_ADC2_IRQ": ("adc 2", "end of conversion"),
        "S32K3_EMIOS0_IRQ": ("emios 0", "interrupt request 23"),
        "S32K3_EMIOS1_IRQ": ("emios 1", "interrupt request 23"),
        "S32K3_EMIOS2_IRQ": ("emios 2", "interrupt request 23"),
        "S32K3_QSPI_IRQ": ("qspi", "tx buffer fill interrupt"),
        "S32K389_GMAC0_IRQ": ("gmac0", "common interrupt"),
        "S32K389_GMAC1_IRQ": ("gmac1", "common interrupt"),
        "S32K3_SAI0_IRQ": ("synchronous audio interface 0", "rx interrupt"),
        "S32K3_SAI1_IRQ": ("synchronous audio interface 1", "rx interrupt"),
        "S32K3_SWT0_IRQ": ("watchdog 0", "platform watchdog initial time out"),
        "S32K3_SWT1_IRQ": ("watchdog 1", "platform watchdog initial time out"),
        "S32K3_SWT2_IRQ": ("watchdog 2", "platform watchdog initial time out"),
        "S32K3_SWT3_IRQ": ("watchdog 3", "platform watchdog initial time out"),
        "S32K3_CONSOLE_LPUART_IRQ": ("lpuart 3", "transmit interrupt"),
    }
    failures: list[str] = []
    for macro, key in expected.items():
        actual = macros.get(macro)
        if actual is None:
            failures.append(f"missing {macro} in {HEADER}")
            continue
        expected_irq = irq_lookup.get((key[0], key[1]))
        if expected_irq is None:
            failures.append(f"spreadsheet row missing for {macro}: {key!r}")
            continue
        if actual != expected_irq:
            failures.append(
                f"{macro}: header {actual} != spreadsheet {expected_irq} "
                f"(expected from {key!r})"
            )
    return failures


def validate_base_macros(macros: dict[str, int], base_lookup: dict[str, int]) -> list[str]:
    expected = {
        "S32K3_ADC0_BASE": "adc 0",
        "S32K3_ADC1_BASE": "adc 1",
        "S32K3_ADC2_BASE": "adc 2",
        "S32K3_EMIOS0_BASE": "emios0",
        "S32K3_EMIOS1_BASE": "emios1",
        "S32K3_EMIOS2_BASE": "emios2",
        "S32K3_QSPI_BASE": "quadspi",
        "S32K3_SAI0_BASE": "sai0",
        "S32K3_SAI1_BASE": "sai1",
        "S32K3_SWT0_BASE": "swt 0",
        "S32K3_SWT1_BASE": "swt 1",
        "S32K3_SWT2_BASE": "swt 2",
        "S32K3_SWT3_BASE": "swt 3",
        "S32K3_LPUART3_BASE": "lpuart 3",
    }
    failures: list[str] = []
    for macro, key_name in expected.items():
        actual = macros.get(macro)
        if actual is None:
            failures.append(f"missing {macro} in {HEADER}")
            continue
        expected_base = base_lookup.get(normalize_name(key_name))
        if expected_base is None:
            failures.append(f"spreadsheet row missing for {macro}: {key_name!r}")
            continue
        if actual != expected_base:
            failures.append(
                f"{macro}: header {actual:#x} != spreadsheet {expected_base:#x}"
            )
    return failures


def main() -> int:
    macros = read_header_macros()
    irq_lookup = build_irq_lookup()
    base_lookup = build_base_lookup()

    failures = []
    failures.extend(validate_irq_macros(macros, irq_lookup))
    failures.extend(validate_base_macros(macros, base_lookup))

    if failures:
        print("S32K389 specification validation failed:")
        for item in failures:
            print(f" - {item}")
        return 1

    print("S32K389 specification validation passed.")
    print(f"Validated IRQ and MMIO definitions from {HEADER.name} against the NXP workbooks.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
